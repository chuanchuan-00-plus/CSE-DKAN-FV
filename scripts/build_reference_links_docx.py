"""Create the clickable reference-link catalogue for the manuscript."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


@dataclass(frozen=True)
class Reference:
    key: str
    citation: str
    url: str
    link_label: str


SECTIONS: list[tuple[str, list[Reference]]] = [
    (
        "DKAN and network components",
        [
            Reference("lei2026dkanpinn", "Lei, G., Exposito, D. & Mao, X. Discontinuity-aware KAN-based physics-informed neural networks. arXiv:2507.08338v2 (2026).", "https://arxiv.org/abs/2507.08338", "arXiv primary record"),
            Reference("liu2024kan", "Liu, Z. et al. KAN: Kolmogorov-Arnold Networks. arXiv:2404.19756 (2024).", "https://arxiv.org/abs/2404.19756", "arXiv primary record"),
            Reference("zhang2025kaf", "Zhang, J., Fan, Y., Cai, K. & Wang, K. Kolmogorov-Arnold Fourier Networks. arXiv:2502.06018 (2025).", "https://arxiv.org/abs/2502.06018", "arXiv primary record"),
            Reference("zhu2025dyt", "Zhu, J., Chen, X., He, K., LeCun, Y. & Liu, Z. Transformers without Normalization. arXiv:2503.10622 (2025).", "https://arxiv.org/abs/2503.10622", "arXiv primary record"),
        ],
    ),
    (
        "Physics-informed learning for shocks",
        [
            Reference("raissi2019pinn", "Raissi, M., Perdikaris, P. & Karniadakis, G. E. Physics-informed neural networks: A deep learning framework for solving forward and inverse problems involving nonlinear partial differential equations. Journal of Computational Physics 378, 686–707 (2019).", "https://doi.org/10.1016/j.jcp.2018.10.045", "DOI record"),
            Reference("wang2021gradient", "Wang, S., Teng, Y. & Perdikaris, P. Understanding and Mitigating Gradient Flow Pathologies in Physics-Informed Neural Networks. SIAM Journal on Scientific Computing 43, A3055–A3081 (2021).", "https://doi.org/10.1137/20M1318043", "DOI record"),
            Reference("krishnapriyan2021failure", "Krishnapriyan, A. S., Gholami, A., Zhe, S., Kirby, R. M. & Mahoney, M. W. Characterizing possible failure modes in physics-informed neural networks. Advances in Neural Information Processing Systems 34 (2021).", "https://proceedings.neurips.cc/paper/2021/hash/df438e5206f31600e6ae4af72f2725f1-Abstract.html", "NeurIPS official record"),
            Reference("patel2022thermodynamic", "Patel, R. G. et al. Thermodynamically consistent physics-informed neural networks for hyperbolic systems. Journal of Computational Physics 449, 110754 (2022).", "https://doi.org/10.1016/j.jcp.2021.110754", "DOI record"),
            Reference("coutinho2023avpinn", "Coutinho, E. J. R. et al. Physics-informed neural networks with adaptive localized artificial viscosity. Journal of Computational Physics 489, 112265 (2023).", "https://doi.org/10.1016/j.jcp.2023.112265", "DOI record"),
            Reference("deryck2024wpinn", "De Ryck, T., Mishra, S. & Molinaro, R. wPINNs: Weak Physics Informed Neural Networks for Approximating Entropy Solutions of Hyperbolic Conservation Laws. SIAM Journal on Numerical Analysis (2024).", "https://doi.org/10.1137/22M1522504", "DOI record"),
        ],
    ),
    (
        "Finite-volume and shock-capturing methods",
        [
            Reference("leveque2002finitevolume", "LeVeque, R. J. Finite Volume Methods for Hyperbolic Problems. Cambridge University Press (2002).", "https://doi.org/10.1017/CBO9780511791253", "DOI record"),
            Reference("toro2009riemann", "Toro, E. F. Riemann Solvers and Numerical Methods for Fluid Dynamics. 3rd edn, Springer (2009).", "https://doi.org/10.1007/b79761", "DOI record"),
            Reference("toro1994hllc", "Toro, E. F., Spruce, M. & Speares, W. Restoration of the contact surface in the HLL-Riemann solver. Shock Waves 4, 25–34 (1994).", "https://doi.org/10.1007/BF01414629", "DOI record"),
            Reference("jiang1996weno", "Jiang, G.-S. & Shu, C.-W. Efficient Implementation of Weighted ENO Schemes. Journal of Computational Physics 126, 202–228 (1996).", "https://doi.org/10.1006/jcph.1996.0130", "DOI record"),
            Reference("borges2008wenoz", "Borges, R., Carmona, M., Costa, B. & Don, W. S. An improved weighted essentially non-oscillatory scheme for hyperbolic conservation laws. Journal of Computational Physics 227, 3191–3211 (2008).", "https://doi.org/10.1016/j.jcp.2007.11.038", "DOI record"),
            Reference("zhang2010positivitydg", "Zhang, X. & Shu, C.-W. On positivity-preserving high order discontinuous Galerkin schemes for compressible Euler equations on rectangular meshes. Journal of Computational Physics 229, 8918–8934 (2010).", "https://doi.org/10.1016/j.jcp.2010.08.016", "DOI record"),
            Reference("zhang2012positivityweno", "Zhang, X. & Shu, C.-W. Positivity-preserving high order finite difference WENO schemes for compressible Euler equations. Journal of Computational Physics 231, 2245–2258 (2012).", "https://doi.org/10.1016/j.jcp.2011.11.020", "DOI record"),
            Reference("sod1978survey", "Sod, G. A. A survey of several finite difference methods for systems of nonlinear hyperbolic conservation laws. Journal of Computational Physics 27, 1–31 (1978).", "https://doi.org/10.1016/0021-9991(78)90023-2", "DOI record"),
        ],
    ),
    (
        "Neural operators and learned discretizations",
        [
            Reference("lu2021deeponet", "Lu, L., Jin, P., Pang, G., Zhang, Z. & Karniadakis, G. E. Learning nonlinear operators via DeepONet based on the universal approximation theorem of operators. Nature Machine Intelligence 3, 218–229 (2021).", "https://doi.org/10.1038/s42256-021-00302-5", "DOI record"),
            Reference("li2021fno", "Li, Z. et al. Fourier Neural Operator for Parametric Partial Differential Equations. International Conference on Learning Representations (2021).", "https://arxiv.org/abs/2010.08895", "arXiv primary record"),
            Reference("takamoto2022pdebench", "Takamoto, M. et al. PDEBench: An Extensive Benchmark for Scientific Machine Learning. Advances in Neural Information Processing Systems, Datasets and Benchmarks Track (2022).", "https://arxiv.org/abs/2210.07182", "arXiv primary record"),
            Reference("barsinai2019discretization", "Bar-Sinai, Y., Hoyer, S., Hickey, J. & Brenner, M. P. Learning data-driven discretizations for partial differential equations. Proceedings of the National Academy of Sciences 116, 15344–15349 (2019).", "https://doi.org/10.1073/pnas.1814058116", "DOI record"),
            Reference("kochkov2021cfd", "Kochkov, D. et al. Machine learning-accelerated computational fluid dynamics. Proceedings of the National Academy of Sciences 118, e2101784118 (2021).", "https://doi.org/10.1073/pnas.2101784118", "DOI record"),
            Reference("lichtle2025unfv", "Lichtlé, N. et al. (U)NFV: Supervised and Unsupervised Neural Finite Volume Methods for Solving Hyperbolic PDEs. arXiv:2505.23702 (2025).", "https://arxiv.org/abs/2505.23702", "arXiv primary record"),
            Reference("cassia2024godunovloss", "Cassia, R. G. & Kerswell, R. R. Godunov Loss Functions for Modelling of Hyperbolic Conservation Laws. arXiv:2405.11674 (2024).", "https://arxiv.org/abs/2405.11674", "arXiv primary record"),
        ],
    ),
]

# Canonical test-problem sources added by the mechanism-generalization revision.
SECTIONS[2][1].extend(
    [
        Reference(
            "lax1954weak",
            "Lax, P. D. Weak Solutions of Nonlinear Hyperbolic Equations and Their Numerical Computation. Communications on Pure and Applied Mathematics 7, 159-193 (1954).",
            "https://doi.org/10.1002/cpa.3160070112",
            "DOI record",
        ),
        Reference(
            "shu1989eno2",
            "Shu, C.-W. & Osher, S. Efficient Implementation of Essentially Non-Oscillatory Shock-Capturing Schemes, II. Journal of Computational Physics 83, 32-78 (1989).",
            "https://doi.org/10.1016/0021-9991(89)90222-2",
            "DOI record",
        ),
        Reference(
            "woodward1984strongshocks",
            "Woodward, P. & Colella, P. The Numerical Simulation of Two-Dimensional Fluid Flow with Strong Shocks. Journal of Computational Physics 54, 115-173 (1984).",
            "https://doi.org/10.1016/0021-9991(84)90142-6",
            "DOI record",
        ),
    ]
)


def clean_citation(text: str) -> str:
    """Remove mojibake left by an earlier Windows code-page conversion."""
    return text.replace("每", "-").replace("谷", "e")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5E87")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.04

    title = document.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(28, 61, 79)

    for name, size in (("Heading 1", 14), ("Heading 2", 11)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(34, 92, 103)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True

    citation = document.styles.add_style("Citation Entry", WD_STYLE_TYPE.PARAGRAPH)
    citation.base_style = document.styles["Normal"]
    citation.font.name = "Aptos"
    citation.font.size = Pt(9.3)
    citation.paragraph_format.left_indent = Cm(0.55)
    citation.paragraph_format.first_line_indent = Cm(-0.55)
    citation.paragraph_format.space_before = Pt(3)
    citation.paragraph_format.space_after = Pt(0)
    citation.paragraph_format.keep_with_next = True

    link = document.styles.add_style("Reference Link", WD_STYLE_TYPE.PARAGRAPH)
    link.base_style = document.styles["Normal"]
    link.font.name = "Aptos"
    link.font.size = Pt(8.6)
    link.paragraph_format.left_indent = Cm(0.55)
    link.paragraph_format.space_after = Pt(4)
    link.paragraph_format.keep_together = True


def build(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    configure_styles(document)

    header = section.header.paragraphs[0]
    header.text = "CSE-DKAN-FV manuscript  |  verified reference links"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(95, 105, 110)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("DKAN + PINN manuscript  |  reference link catalogue")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(95, 105, 110)

    title = document.add_paragraph(style="Title")
    title.add_run("Reference Link Catalogue")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("Conservative Shock-Explicit DKAN Finite-Volume Learning")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(49, 80, 91)

    note = document.add_paragraph()
    note.add_run("Purpose. ").bold = True
    note.add_run(
        "Use the links below to open the publisher or arXiv record, then export BibTeX into Overleaf. "
        "The supplied LaTeX package already contains a compile-ready references.bib file. Metadata was checked on 20 July 2026."
    )

    document.add_heading("BibTeX import workflow", level=1)
    for text in (
        "Open the primary link and choose Cite, Export citation, or BibTeX.",
        "Keep the suggested key shown in square brackets, or update the matching key in main.tex.",
        "After importing into Overleaf, compile twice and check that no citation is undefined.",
    ):
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(text)

    number = 1
    for heading, references in SECTIONS:
        document.add_heading(heading, level=1)
        for reference in references:
            paragraph = document.add_paragraph(style="Citation Entry")
            key_run = paragraph.add_run(f"{number}. [{reference.key}] ")
            key_run.bold = True
            paragraph.add_run(clean_citation(reference.citation))
            link_paragraph = document.add_paragraph(style="Reference Link")
            link_paragraph.add_run(f"{reference.link_label}: ")
            add_hyperlink(link_paragraph, reference.url, reference.url)
            number += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    build(Path("paper_build/reference_links/DKAN_PINN_reference_links.docx"))
